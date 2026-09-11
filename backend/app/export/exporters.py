import os
import re
import uuid

import pandas as pd
from docx import Document as DocxDocument
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from app.config import settings
from app.models import Document, Node

EXPORT_DIR = os.path.join(settings.upload_dir, "..", "exports")


def _ensure_export_dir() -> str:
    path = os.path.abspath(EXPORT_DIR)
    os.makedirs(path, exist_ok=True)
    return path


def _safe_sheet_name(name: str) -> str:
    """Excel sheet names: max 31 chars, no []:*?/\\ characters."""
    cleaned = re.sub(r"[\[\]\:\*\?/\\]", "_", name)
    return cleaned[:31] if cleaned else "Sheet"


def export_document_to_docx(doc: Document, nodes: list[Node]) -> str:
    out_dir = _ensure_export_dir()
    docx = DocxDocument()
    docx.add_heading(doc.writer_name or doc.student_id, level=1)
    docx.add_paragraph(f"Student ID: {doc.student_id}")
    if doc.class_id:
        docx.add_paragraph(f"Class: {doc.class_id}")
    docx.add_paragraph(f"Source file: {doc.filename}")
    docx.add_paragraph("")

    for node in nodes:
        p = docx.add_paragraph(node.corrected_text)
        if node.confidence_score < settings.low_confidence_threshold:
            p.add_run("  [low confidence — please review]").italic = True

    filepath = os.path.join(out_dir, f"{doc.student_id}-{uuid.uuid4().hex[:8]}.docx")
    docx.save(filepath)
    return filepath


def export_document_to_pdf(doc: Document, nodes: list[Node]) -> str:
    out_dir = _ensure_export_dir()
    filepath = os.path.join(out_dir, f"{doc.student_id}-{uuid.uuid4().hex[:8]}.pdf")

    styles = getSampleStyleSheet()
    pdf = SimpleDocTemplate(filepath, pagesize=LETTER)
    story = [
        Paragraph(doc.writer_name or doc.student_id, styles["Title"]),
        Paragraph(f"Student ID: {doc.student_id}", styles["Normal"]),
    ]
    if doc.class_id:
        story.append(Paragraph(f"Class: {doc.class_id}", styles["Normal"]))
    story.append(Spacer(1, 16))

    for node in nodes:
        text = node.corrected_text
        if node.confidence_score < settings.low_confidence_threshold:
            text += " <i>[low confidence]</i>"
        story.append(Paragraph(text, styles["BodyText"]))
        story.append(Spacer(1, 6))

    pdf.build(story)
    return filepath


def export_class_to_xlsx(docs: list[Document], nodes_by_doc: dict[uuid.UUID, list[Node]]) -> str:
    """One workbook, one tab per student. Each tab lists that student's
    recognized lines with raw/corrected text and confidence, so a teacher
    can scan a whole class batch at once."""
    out_dir = _ensure_export_dir()
    class_label = docs[0].class_id or "export"
    filepath = os.path.join(out_dir, f"class-{class_label}-{uuid.uuid4().hex[:8]}.xlsx")

    used_names: set[str] = set()
    with pd.ExcelWriter(filepath, engine="openpyxl") as writer:
        for doc in docs:
            nodes = nodes_by_doc.get(doc.id, [])
            rows = [
                {
                    "raw_text": n.raw_text,
                    "corrected_text": n.corrected_text,
                    "was_corrected": n.is_corrected,
                    "confidence": round(n.confidence_score, 3),
                    "needs_review": n.confidence_score < settings.low_confidence_threshold,
                }
                for n in nodes
            ]
            df = pd.DataFrame(rows) if rows else pd.DataFrame(
                columns=["raw_text", "corrected_text", "was_corrected", "confidence", "needs_review"]
            )

            base_name = _safe_sheet_name(doc.writer_name or doc.student_id)
            sheet_name = base_name
            suffix = 1
            while sheet_name in used_names:
                suffix += 1
                sheet_name = _safe_sheet_name(f"{base_name}_{suffix}")
            used_names.add(sheet_name)

            df.to_excel(writer, sheet_name=sheet_name, index=False)

    return filepath
